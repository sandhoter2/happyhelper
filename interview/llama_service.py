"""Local document text extraction — replaces LlamaCloud with pypdf + python-docx."""
import os

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".txt", ".xlsx", ".xls",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".go",
    ".rb", ".rs", ".cpp", ".c", ".h", ".cs", ".md",
}


def extract_text(file_obj, filename: str) -> str:
    """Extract plain text from a file object. Returns text string."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_obj)
    if ext == ".docx":
        return _extract_docx(file_obj)
    if ext in (".xlsx", ".xls"):
        return _extract_excel(file_obj)
    # .doc, .txt, .md, .py, .js, .ts … treat as UTF-8 text
    try:
        raw = file_obj.read()
        return raw.decode("utf-8", errors="replace")
    except Exception:
        return ""


def _extract_pdf(file_obj) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(file_obj)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text and text.strip():
                pages.append(text.strip())
        return "\n\n".join(pages)
    except Exception:
        return ""


def _extract_docx(file_obj) -> str:
    try:
        from docx import Document
        doc = Document(file_obj)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception:
        return ""


def _extract_excel(file_obj) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = "\t".join("" if c is None else str(c) for c in row)
                if line.strip():
                    rows.append(line)
        return "\n".join(rows)
    except Exception:
        return ""


def estimate_page_count(text: str) -> int:
    if not text:
        return 0
    return max(1, len(text) // 3000)


# ── Stub interface — kept so views.py callers need minimal changes ────────

def upload_file(file_obj, filename: str, webhook_id: str = "") -> dict:
    """Extract text synchronously and return it in the result dict."""
    text = extract_text(file_obj, filename)
    return {"id": "", "status": "COMPLETED", "file_id": "", "_text": text}


def get_job_status(job_id: str) -> dict:
    return {"status": "COMPLETED"}


def get_parsed_markdown(job_id: str) -> str:
    return ""


def delete_job(job_id: str) -> bool:
    return True


def delete_file(file_id: str) -> bool:
    return True


def poll_until_done(job_id: str, max_wait: int = 120) -> str:
    return "COMPLETED"


# ── Cloud directory / index stubs (no-ops) ───────────────────────────────

def create_directory(name: str) -> dict:
    return {"id": "", "name": name}


def add_file_to_directory(directory_id: str, file_id: str) -> dict:
    return {}


def remove_file_from_directory(directory_id: str, file_id: str) -> bool:
    return True


def delete_directory(directory_id: str) -> bool:
    return True


def create_index(directory_id: str, name: str) -> dict:
    return {"id": "", "name": name}


def query_index(index_id: str, query: str, top_k: int = 5) -> list:
    return []


def delete_index(index_id: str) -> bool:
    return True


def get_file_status(file_id: str) -> dict:
    return {"id": file_id, "status": "available"}


def list_directory_files(directory_id: str) -> list:
    return []
